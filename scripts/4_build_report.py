"""
Step 4 — Build the multi-sheet Excel report.

Reads:   data/rows.json, data/cost.json, data/cost_window.json,
         data/vms.json, data/submap.json, config.json
Writes:  <report>.xlsx  (name derived from config.reportTitle, or $OUTNAME)

Run from the repository root:  python scripts/4_build_report.py
"""

import collections
import datetime
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import classify as cls

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, 'data')


def load_json(path, default=None):
    if default is not None and not os.path.exists(path):
        return default
    with open(path, encoding='utf-8') as f:
        return json.load(f)


config = load_json(os.path.join(ROOT, 'config.json'))
rows_data = load_json(os.path.join(DATA, 'rows.json'))
retire_rows = rows_data['retire_rows']
price_rows = rows_data['price_rows']
window = load_json(os.path.join(DATA, 'cost_window.json'))
cost = load_json(os.path.join(DATA, 'cost.json'), default=[])
vms = load_json(os.path.join(DATA, 'vms.json'), default=[])
submap = load_json(os.path.join(DATA, 'submap.json'), default={})

PCT = float(config.get('priceIncreasePercent', 25)) / 100.0
WINDOW_DAYS = int(config.get('costWindowDays', 30))
ORG = config.get('organizationName', 'Organization')
TITLE = config.get('reportTitle', 'Azure VM Retirement & Price-Change Impact')
ADVISORY = config.get('advisoryReference', 'Microsoft Health Advisory')
TENANT = config.get('tenantId', '')
EOL = config.get('retirementEndOfLife', '')
EFFECTIVE = config.get('priceIncreaseEffective', '')
RETIRE_FAMS = config.get('retirementFamilies', [])
V1_FAMS = config.get('priceFamiliesV1', [])
V2_FAMS = config.get('priceFamiliesV2', [])
TAG_HEADERS = [c['header'] for c in config.get('tagColumns', [])]

SUB_COUNT = len(submap)
VM_COUNT = len(vms)

# ---------------- cost aggregation by family group ----------------
cost_group = collections.defaultdict(float)   # (wave, label) -> cost
wave_total = collections.defaultdict(float)
for r in cost:
    label, wave = cls.meter_group(r.get('Meter') or '')
    if label:
        c = r.get('Cost') or 0
        cost_group[(wave, label)] += c
        wave_total[wave] += c

v1_total = wave_total['v1']
v2_total = wave_total['v2']
grand = v1_total + v2_total

retire_ct = collections.Counter(r['Family'] for r in retire_rows)
price_ct = collections.Counter(r['Family'] for r in price_rows)

# ---------------- workbook styling ----------------
HDR_FILL = PatternFill('solid', fgColor='1F4E78')
HDR_FONT = Font(bold=True, color='FFFFFF', size=11)
TITLE_FONT = Font(bold=True, size=16, color='1F4E78')
SUB_FONT = Font(bold=True, size=12, color='1F4E78')
NOTE_FONT = Font(italic=True, size=9, color='555555')
BOLD = Font(bold=True)
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00'


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER


def write_table(ws, start_row, headers, rows, widths=None, money_cols=()):
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    style_header(ws, start_row, len(headers))
    r = start_row + 1
    for rowd in rows:
        for j, h in enumerate(headers, 1):
            val = rowd.get(h, '')
            cell = ws.cell(row=r, column=j, value=val)
            cell.border = BORDER
            if j in money_cols and isinstance(val, (int, float)):
                cell.number_format = MONEY
        r += 1
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f'A{start_row}:{last_col}{r - 1}'
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return r


wb = Workbook()

# ============ Sheet 1: Summary ============
ws = wb.active
ws.title = 'Summary'
ws.sheet_view.showGridLines = False
ws['A1'] = f'{ORG} — {TITLE}'
ws['A1'].font = TITLE_FONT
ws['A2'] = f"Tenant {TENANT}  |  {SUB_COUNT:,} enabled subscriptions  |  {VM_COUNT:,} VMs scanned"
ws['A2'].font = NOTE_FONT
cost_from = window['from'][:10] if window else ''
cost_to = window['to'][:10] if window else ''
ws['A3'] = (f"Source: {ADVISORY}  |  Generated {datetime.date.today().isoformat()}  |  "
            f"Cost window {cost_from} to {cost_to} ({WINDOW_DAYS} days, AmortizedCost)")
ws['A3'].font = NOTE_FONT

r = 5
ws.cell(row=r, column=1,
        value=f"1 - Retirements (End of Life {EOL}): {', '.join(RETIRE_FAMS)}").font = SUB_FONT
r += 1
for j, h in enumerate(['Family', 'VM Count'], 1):
    ws.cell(row=r, column=j, value=h)
style_header(ws, r, 2)
r += 1
for fam in RETIRE_FAMS:
    ws.cell(row=r, column=1, value=fam).border = BORDER
    ws.cell(row=r, column=2, value=retire_ct.get(fam, 0)).border = BORDER
    r += 1
ws.cell(row=r, column=1, value='TOTAL').font = BOLD
ws.cell(row=r, column=1).border = BORDER
tc = ws.cell(row=r, column=2, value=sum(retire_ct.values()))
tc.font = BOLD
tc.border = BORDER
r += 2

pct_label = f"~{int(round(PCT * 100))}%"
ws.cell(row=r, column=1,
        value=f"2 - Price Increase {pct_label} (effective {EFFECTIVE}): v1 & v2 series").font = SUB_FONT
r += 1
for j, h in enumerate(['Change Wave', 'Family', 'VM Count'], 1):
    ws.cell(row=r, column=j, value=h)
style_header(ws, r, 3)
r += 1
for wave, fams in [('v1', V1_FAMS), ('v2', V2_FAMS)]:
    for fam in fams:
        if price_ct.get(fam, 0):
            ws.cell(row=r, column=1, value=wave).border = BORDER
            ws.cell(row=r, column=2, value=fam).border = BORDER
            ws.cell(row=r, column=3, value=price_ct[fam]).border = BORDER
            r += 1
ws.cell(row=r, column=1, value='TOTAL').font = BOLD
ws.cell(row=r, column=1).border = BORDER
ws.cell(row=r, column=2, value='').border = BORDER
tc = ws.cell(row=r, column=3, value=sum(price_ct.values()))
tc.font = BOLD
tc.border = BORDER
r += 2

ws.cell(row=r, column=1,
        value=f"3 - Estimated cost impact of +{pct_label} (amortized, trailing {WINDOW_DAYS} days)").font = SUB_FONT
r += 1
for j, h in enumerate(['Change Wave', f'Amortized {WINDOW_DAYS}d', f'+{pct_label} Increase',
                       f'Projected {WINDOW_DAYS}d', 'Projected Annualized'], 1):
    ws.cell(row=r, column=j, value=h)
style_header(ws, r, 5)
r += 1
for wave, label, tot in [('v1', 'v1 series', v1_total),
                         ('v2', 'v2 series', v2_total),
                         ('TOTAL', 'All impacted', grand)]:
    ws.cell(row=r, column=1, value=label).border = BORDER
    annual_factor = 365.0 / WINDOW_DAYS
    for col, val in [(2, tot), (3, tot * PCT), (4, tot * (1 + PCT)),
                     (5, tot * (1 + PCT) * annual_factor)]:
        cc = ws.cell(row=r, column=col, value=round(val, 2))
        cc.number_format = MONEY
        cc.border = BORDER
        if wave == 'TOTAL':
            cc.font = BOLD
    if wave == 'TOTAL':
        ws.cell(row=r, column=1).font = BOLD
    r += 1
r += 1

notes = [
    'Notes:',
    f'- Counts and inventories are tenant-wide across all {SUB_COUNT:,} enabled subscriptions (Azure Resource Graph).',
    f'- Cost = AmortizedCost for MeterCategory "Virtual Machines" over the trailing {WINDOW_DAYS} days, '
    f'summed for impacted v1/v2 SKUs, then +{pct_label}.',
    '- Price increase applies to PAYG, EA, and MCA rates. It does NOT change existing Reserved Instances '
    '(RI pricing is locked at purchase);',
    '   amortized cost includes RI-covered usage, so the increase figure is an upper bound for any RI-covered VMs.',
    f'- {", ".join(RETIRE_FAMS)} are being RETIRED (not price-increased) and are excluded from the cost estimate.',
    '- Azure bills some sizes on combined meters (e.g., "D4 v2/DS4 v2"); cost is grouped by billing meter series accordingly.',
]
for n in notes:
    ws.cell(row=r, column=1, value=n).font = NOTE_FONT
    r += 1

ws.column_dimensions['A'].width = 26
for col in ['B', 'C', 'D', 'E']:
    ws.column_dimensions[col].width = 20

# ---------------- inventory sheet columns ----------------
BASE_COLS = ['Subscription ID', 'Subscription Name', 'Resource Group', 'VM Name', 'VM Size', 'Family']
TAIL_COLS = ['Power State', 'All Tags']
BASE_WIDTHS = [38, 32, 30, 32, 22, 8]
TAG_WIDTH = 20
TAIL_WIDTHS = [16, 70]

# ============ Sheet 2: Retirement Inventory ============
ws2 = wb.create_sheet('Retirement Inventory')
headers2 = BASE_COLS + TAG_HEADERS + TAIL_COLS
widths2 = BASE_WIDTHS + [TAG_WIDTH] * len(TAG_HEADERS) + TAIL_WIDTHS
retire_sorted = sorted(retire_rows, key=lambda x: (x['Subscription Name'], x['Resource Group'], x['VM Name']))
write_table(ws2, 1, headers2, retire_sorted, widths2)

# ============ Sheet 3: Price Increase Inventory ============
ws3 = wb.create_sheet('Price Increase Inventory')
headers3 = BASE_COLS + ['Change Wave'] + TAG_HEADERS + TAIL_COLS
widths3 = BASE_WIDTHS + [12] + [TAG_WIDTH] * len(TAG_HEADERS) + TAIL_WIDTHS
price_sorted = sorted(price_rows, key=lambda x: (x['Change Wave'], x['Subscription Name'],
                                                 x['Resource Group'], x['VM Name']))
write_table(ws3, 1, headers3, price_sorted, widths3)

# ============ Sheet 4: Cost Impact ============
ws4 = wb.create_sheet('Cost Impact')
ws4['A1'] = f'Estimated +{pct_label} Price Impact - Amortized VM cost, trailing {WINDOW_DAYS} days'
ws4['A1'].font = SUB_FONT
ws4['A2'] = f"Window {cost_from} to {cost_to}  |  MeterCategory = Virtual Machines"
ws4['A2'].font = NOTE_FONT
chdr = ['Change Wave', 'Family / Meter Group', f'Amortized Cost ({WINDOW_DAYS}d)',
        f'+{pct_label} Increase', f'Projected Cost ({WINDOW_DAYS}d)']
cost_rows_out = []
for (wave, label), c in sorted(cost_group.items(), key=lambda kv: (kv[0][0], -kv[1])):
    cost_rows_out.append({
        'Change Wave': wave,
        'Family / Meter Group': label,
        f'Amortized Cost ({WINDOW_DAYS}d)': round(c, 2),
        f'+{pct_label} Increase': round(c * PCT, 2),
        f'Projected Cost ({WINDOW_DAYS}d)': round(c * (1 + PCT), 2),
    })
endr = write_table(ws4, 4, chdr, cost_rows_out, [14, 26, 22, 18, 20], money_cols=(3, 4, 5))
ws4.cell(row=endr, column=1, value='TOTAL').font = BOLD
ws4.cell(row=endr, column=2, value='All impacted v1+v2').font = BOLD
for col, val in [(3, grand), (4, grand * PCT), (5, grand * (1 + PCT))]:
    cc = ws4.cell(row=endr, column=col, value=round(val, 2))
    cc.number_format = MONEY
    cc.font = BOLD
    cc.border = BORDER
for c in range(1, 6):
    ws4.cell(row=endr, column=c).border = BORDER

# ---------------- save ----------------
default_name = re.sub(r'[^A-Za-z0-9]+', '_', f'{ORG}_{TITLE}').strip('_') + '.xlsx'
out = os.environ.get('OUTNAME') or os.path.join(ROOT, default_name)
wb.save(out)
print('Saved', out)
print('Retirement rows:', len(retire_rows), '| Price rows:', len(price_rows))
print('v1 cost ${:,.2f} | v2 cost ${:,.2f} | total ${:,.2f} | +{:.0f}% ${:,.2f}'.format(
    v1_total, v2_total, grand, PCT * 100, grand * PCT))
print('Cost groups:')
for (wave, label), c in sorted(cost_group.items()):
    print(f'  {wave}  {label:14} ${c:,.2f}')
